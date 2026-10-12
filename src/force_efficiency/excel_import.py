from __future__ import annotations

import hashlib
import json
import re
from dataclasses import replace
from pathlib import Path
from platform_config import project_path

from .models import GroupCandidate, MeasurementCandidate


DEFAULT_SOURCE_FILES = (
    project_path("force_raw") / "7寸桨叶.xlsx",
    project_path("force_raw") / "7寸桨叶-20260121.xlsx",
)
EXCLUDED_SHEETS = {"校准"}
EXCLUDED_GROUPS = {
    ("7寸桨叶.xlsx", "1806-2280kv", "B21"),
    ("7寸桨叶-20260121.xlsx", "2205-1700kV", "B3"),
    ("7寸桨叶-20260121.xlsx", "2205-1700kV", "B9"),
    ("7寸桨叶-20260121.xlsx", "1806-1400kV", "B3"),
    ("7寸桨叶-20260121.xlsx", "1806-1400kV", "B9"),
}
SKIPPED_WORKBOOK_SHEETS = {("7寸桨叶-20260121.xlsx", "2205-1700kV-0118")}
PROPELLER_ALIASES = {
    "7040-3hsrc": "HSRC 7040-3",
    "7035-2gf": "GF 7035-2",
    "7042-2gf": "GF 7042-2",
    "7036f-2gf": "GF 7036F-2",
}
TEST_DATE_OVERRIDES = {
    ("7寸桨叶.xlsx", "2306-1850kv", "V3"): "1月14日",
    ("7寸桨叶.xlsx", "2306-1850kv", "V12"): "1月14日",
    ("7寸桨叶.xlsx", "2306-1850kv", "AF9"): "1月14日",
    ("7寸桨叶.xlsx", "2306-1850kv", "AF18"): "1月14日",
    ("7寸桨叶.xlsx", "2306-1850kV-0117", "*"): "1月17日",
    ("7寸桨叶-20260121.xlsx", "2205-1700kV-0122", "*"): "1月22日",
}
THROTTLE_OVERRIDES = {
    ("7寸桨叶.xlsx", "2205-1700kV-0118", "B3", 5): 1900.0,
    ("7寸桨叶-20260121.xlsx", "2205-1700kV-0118", "B3", 5): 1900.0,
}


def number(value: object) -> float | None:
    return float(value) if isinstance(value, (int, float)) and not isinstance(value, bool) else None


def number_in_text(value: object) -> float | None:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    match = re.search(r"[-+]?\d+(?:\.\d+)?", str(value)) if value is not None else None
    return float(match.group()) if match else None


def motor_spec(sheet_name: str, sheet: object, throttle_column: int, header_row: int) -> str:
    if sheet_name.strip().casefold() == "1506-3000kv":
        return "致盈 1506-3000kV"
    if sheet_name.strip().casefold() == "2306-1850kv":
        return "丰云 2306-1850kV"
    if sheet_name.strip().casefold() == "2306-1850kv-0117":
        return "丰云 2306-1850kV"
    if sheet_name.strip().casefold() == "2004-2850kv":
        return "AMAX 2004-2850kV"
    if sheet_name.strip().casefold() == "2004-2400kv":
        return "AMAX 2004-2400kV"
    if sheet_name.strip().casefold() == "1806-2280kv":
        return "AOSENMA 1806-2280kV"
    if sheet_name.strip().casefold() == "2205-1700kv":
        return "HSKRC 2205-1700kV"
    if sheet_name.strip().casefold() == "2205-1700kv-0118":
        return "HSKRC 2205-1700kV"
    if sheet_name.strip().casefold() == "2205-1700kv-0122":
        return "HSKRC 2205-1700kV"
    if sheet_name.strip().casefold() == "2205-1700kv-重-0122":
        return "HSKRC 2205 1850kV"
    if sheet_name.strip().casefold() == "2306.5-1500kv":
        return "丰云2306.5-1500kV"
    if sheet_name.strip().casefold() == "2005-1500kv":
        return "AMAX 2005-1500kV"
    if sheet_name.strip().casefold() == "1806-1400kv":
        return "东兴威 1806-1400kV"
    if sheet_name.strip().casefold() == "2810-1350kv":
        return "某出口 2810-1350kV"
    base = re.sub(r"-(?:重-)?\d{4}$", "", sheet_name, flags=re.IGNORECASE)
    rated_kv = None
    for row in range(max(1, header_row - 4), header_row):
        if str(sheet.cell(row, throttle_column).value).strip().upper() == "KV":
            rated_kv = number_in_text(sheet.cell(row, throttle_column + 1).value)
    size = re.match(r"\d+(?:\.\d+)?", base)
    if size and rated_kv is not None:
        return f"{size.group()}-{rated_kv:g}kV"
    return re.sub(r"kv", "kV", base, flags=re.IGNORECASE)


def propeller_name(value: object) -> str | None:
    text = str(value).strip() if value is not None else ""
    if not text:
        return None
    return PROPELLER_ALIASES.get(text.casefold(), text)


def test_date(path: Path, sheet_name: str, anchor: str) -> str | None:
    return TEST_DATE_OVERRIDES.get(
        (path.name, sheet_name, anchor),
        TEST_DATE_OVERRIDES.get((path.name, sheet_name, "*")),
    )


def split_by_voltage(group: GroupCandidate) -> list[GroupCandidate]:
    """把同一数据块内不同实测电压拆成独立数据组，保持首次出现顺序。"""
    buckets: dict[float | None, list[MeasurementCandidate]] = {}
    order: list[float | None] = []
    for point in group.points:
        voltage = round(point.voltage_v, 9) if point.voltage_v is not None else None
        if voltage not in buckets:
            buckets[voltage] = []
            order.append(voltage)
        buckets[voltage].append(point)
    if len(order) <= 1:
        return [group]
    return [
        replace(
            group,
            nominal_voltage_v=voltage,
            points=tuple(replace(point, point_order=index) for index, point in enumerate(buckets[voltage], 1)),
        )
        for voltage in order
    ]


def scan_workbooks(paths: tuple[Path, ...] | list[Path]) -> list[GroupCandidate]:
    """只读扫描；顺序为工作簿、工作表、数据块列、数据块行。"""
    import openpyxl

    groups: list[GroupCandidate] = []
    source_order = 1
    for workbook_index, path in enumerate(paths, 1):
        formulas = openpyxl.load_workbook(path, data_only=False, read_only=False)
        values = openpyxl.load_workbook(path, data_only=True, read_only=False)
        try:
            for sheet_index, (formula_sheet, value_sheet) in enumerate(
                zip(formulas.worksheets, values.worksheets), 1
            ):
                if formula_sheet.title.strip() in EXCLUDED_SHEETS:
                    continue
                skipped_sheet = (path.name, formula_sheet.title) in SKIPPED_WORKBOOK_SHEETS
                anchors = sorted(
                    (
                        cell
                        for row in formula_sheet.iter_rows()
                        for cell in row
                        if cell.value == "油门"
                    ),
                    key=lambda cell: (cell.column, cell.row),
                )
                for block_index, anchor in enumerate(anchors, 1):
                    candidate = parse_group(
                        formula_sheet,
                        value_sheet,
                        path,
                        workbook_index,
                        sheet_index,
                        block_index,
                        anchor,
                        source_order,
                    )
                    excluded = (path.name, formula_sheet.title, anchor.coordinate) in EXCLUDED_GROUPS
                    for split_group in split_by_voltage(candidate):
                        assigned_order = source_order
                        source_order += 1
                        if excluded or skipped_sheet:
                            continue
                        if split_group.points:
                            groups.append(replace(split_group, source_order=assigned_order))
        finally:
            formulas.close()
            values.close()
    return groups


def parse_group(
    formula_sheet: object,
    value_sheet: object,
    path: Path,
    workbook_index: int,
    sheet_index: int,
    block_index: int,
    anchor: object,
    source_order: int,
) -> GroupCandidate:
    header_row, throttle_column = anchor.row, anchor.column
    no_load_row = next(
        (
            row
            for row in range(header_row - 1, max(0, header_row - 4), -1)
            if formula_sheet.cell(row, throttle_column).value == "空载电流"
        ),
        None,
    )
    settings_row = no_load_row - 1 if no_load_row and no_load_row > 1 else None
    no_load_current = number(value_sheet.cell(no_load_row, throttle_column + 1).value) if no_load_row else None
    propeller_value = formula_sheet.cell(no_load_row, throttle_column + 2).value if no_load_row else None
    propeller = propeller_name(propeller_value)

    advance = number_in_text(formula_sheet.cell(settings_row, throttle_column).value) if settings_row else None
    frequency = number_in_text(formula_sheet.cell(settings_row, throttle_column + 1).value) if settings_row else None
    configured_kv = number_in_text(formula_sheet.cell(settings_row, throttle_column + 2).value) if settings_row else None

    points: list[MeasurementCandidate] = []
    row = header_row + 1
    while row <= formula_sheet.max_row and number(formula_sheet.cell(row, throttle_column).value) is not None:
        throttle, voltage, current, rpm, thrust_mn, power, thrust_g, efficiency = (
            number(value_sheet.cell(row, throttle_column + offset).value) for offset in range(8)
        )
        throttle = THROTTLE_OVERRIDES.get(
            (path.name, formula_sheet.title, anchor.coordinate, row), throttle
        )
        # Derived fields are deliberately recomputed; Excel formula caches may reference another block.
        power = (
            voltage * (current - no_load_current)
            if voltage is not None and current is not None and no_load_current is not None
            else None
        )
        thrust_g = thrust_mn / 9.8 if thrust_mn is not None else None
        efficiency = thrust_g / power if thrust_g is not None and power not in (None, 0) else None
        points.append(
            MeasurementCandidate(
                len(points) + 1,
                throttle,
                voltage,
                current,
                rpm,
                thrust_mn,
                power,
                thrust_g,
                efficiency,
                row,
            )
        )
        row += 1

    nominal_voltage = next((point.voltage_v for point in points if point.voltage_v is not None), None)
    return GroupCandidate(
        source_order=source_order,
        motor_spec=motor_spec(formula_sheet.title, formula_sheet, throttle_column, header_row),
        test_date=test_date(path, formula_sheet.title, anchor.coordinate),
        propeller=propeller,
        nominal_voltage_v=nominal_voltage,
        esc_advance_deg=advance,
        pwm_frequency_khz=frequency,
        am32_motor_kv=round(configured_kv) if configured_kv is not None else None,
        no_load_current_a=no_load_current,
        source_file=path.name,
        source_workbook_index=workbook_index,
        source_sheet=formula_sheet.title,
        source_sheet_index=sheet_index,
        source_block_index=block_index,
        source_anchor=anchor.coordinate,
        points=tuple(points),
    )


def fingerprint(group: GroupCandidate) -> str:
    """用设置参数及原始测量列找重复；不依赖 Excel 公式缓存值。"""
    def normalized(value: object) -> object:
        if isinstance(value, float):
            return round(value, 6)
        return value.strip().casefold() if isinstance(value, str) else value

    payload = {
        "settings": [normalized(value) for value in (
            group.motor_spec,
            group.propeller,
            group.nominal_voltage_v,
            group.esc_advance_deg,
            group.pwm_frequency_khz,
            group.am32_motor_kv,
            group.no_load_current_a,
        )],
        "points": [
            [normalized(value) for value in (
                point.throttle_us, point.voltage_v, point.current_a, point.rpm, point.thrust_mn
            )]
            for point in group.points
        ],
    }
    return hashlib.sha256(json.dumps(payload, ensure_ascii=False, sort_keys=True).encode("utf-8")).hexdigest()


def duplicate_candidates(groups: list[GroupCandidate]) -> list[list[GroupCandidate]]:
    buckets: dict[str, list[GroupCandidate]] = {}
    for group in groups:
        buckets.setdefault(fingerprint(group), []).append(group)
    return [bucket for bucket in buckets.values() if len(bucket) > 1]


def candidate_name(group: GroupCandidate) -> str:
    voltage = "?V" if group.nominal_voltage_v is None else f"{group.nominal_voltage_v:g}V"
    return f"{group.source_order:03d}-{group.motor_spec}-{group.propeller or '桨叶待确认'}-{voltage}"


def import_preview(paths: tuple[Path, ...] | list[Path]) -> dict[str, object]:
    groups = scan_workbooks(paths)
    return {
        "group_count": len(groups),
        "point_count": sum(len(group.points) for group in groups),
        "ordering": "工作簿 -> 工作表 -> 数据块列从左到右 -> 同列从上到下",
        "excluded_sheets": sorted(EXCLUDED_SHEETS),
        "groups": [
            {
                "source_order": group.source_order,
                "name": candidate_name(group),
                "test_date": group.test_date,
                "source": f"{group.source_file}/{group.source_sheet}/{group.source_anchor}",
                "points": len(group.points),
            }
            for group in groups
        ],
        "duplicate_candidates": [
            [
                {
                    "source_order": group.source_order,
                    "name": candidate_name(group),
                    "source": f"{group.source_file}/{group.source_sheet}/{group.source_anchor}",
                }
                for group in bucket
            ]
            for bucket in duplicate_candidates(groups)
        ],
    }
