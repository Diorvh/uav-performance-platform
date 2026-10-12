"""Read-only workbook previews. No database dependency or workbook writes."""
from __future__ import annotations
from datetime import datetime, date, timedelta
from pathlib import Path
import hashlib
import json
import re
import statistics
import csv
from contextlib import contextmanager
import openpyxl
from openpyxl.utils import get_column_letter
from .calculations import calculate_points


def text(value):
    if value is None:
        return ''
    return str(int(value)) if isinstance(value, float) and value.is_integer() else str(value).strip()


def form_factor(path):
    for name in ('21700', '18650'):
        if name in str(path):
            return name
    if '软包' in str(path):
        return '软包'
    return '其他' if Path(path).name == '其他电池记录.xlsx' else ''


def catalog_value(value):
    return value.isoformat(' ') if isinstance(value, datetime) else value.isoformat() if isinstance(value, date) else value


def catalog_cells(row, rownum, header_rows):
    """Retain original column identity even for duplicate or unnamed headers."""
    result = []
    for index, value in enumerate(row):
        context = [text(h[index]) for h in header_rows if index < len(h) and h[index] is not None]
        if value is None and not context:
            continue
        column = get_column_letter(index + 1)
        result.append(dict(column=column, cell=f'{column}{rownum}',
            label=context[-1] if context else f'{column}列（原表未命名）',
            header_context=context, value=catalog_value(value)))
    return result


def catalog_tests(row, rownum, start, headers):
    """Extract reference tests only; never create a discharge experiment."""
    result = []
    for index in range(start, len(row)):
        value = row[index]
        if not isinstance(value, (date, datetime)):
            continue
        tail = list(row[index+1:index+6]) + [None] * 5
        unit_header = text(headers[index+3]) if index+3 < len(headers) else ''
        unit = 'Wh' if unit_header.lower() == 'wh' else 'mWh' if unit_header.lower() == 'mwh' else unit_header
        energy_note = ''
        if unit == 'mWh' and isinstance(tail[2], (int, float)) and isinstance(tail[1], (int, float)) and tail[1] > 0 and 0 < tail[2] / tail[1] < .1:
            energy_note = '原表能量单位与数值可能不一致，保留原值待核实'
        result.append(dict(source_row=rownum, source_cell=f'{get_column_letter(index+1)}{rownum}',
            date=catalog_value(value), device=tail[0], capacity_mah=tail[1],
            energy_value=tail[2], energy_unit=unit, energy_note=energy_note,
            energy_mwh=tail[2] if unit == 'mWh' else None,
            energy_wh=tail[2] if unit == 'Wh' else None, conditions=tail[3],
            charge_conditions=tail[4] if index+5 < len(headers) and text(headers[index+5]) == '充电' else None))
    return result


def catalog_identity(row):
    code = row['code']
    if row['form_factor'] == '软包':
        # Serial numbers are identifiers, not cylinder category/sample pairs.
        category = row.get('model_spec') or row.get('appearance') or '型号未确认'
        if row.get('nominal_voltage_v') is not None:
            category += f"-{row['nominal_voltage_v']:g}V"
        if row.get('nominal_capacity_ah') is not None:
            category += f"-{row['nominal_capacity_ah']*1000:g}mAh"
        return category, code
    return tuple(code.split('-', 1)) if '-' in code else (code, '原编号'+code)


def load_mappings(paths):
    records, issues = [], []
    for path in paths:
        try:
            source_hash = hashlib.sha256(Path(path).read_bytes()).hexdigest()
            book = openpyxl.load_workbook(path, read_only=True, data_only=True)
            try:
                for sheet in book:
                    headers = None
                    header_rows = []
                    current_record = None
                    for rownum, row in enumerate(sheet.iter_rows(values_only=True), 1):
                        values = [text(v) for v in row]
                        if headers is None:
                            header_rows.append(row)
                            if '编号' in values and '厂家' in values:
                                headers = {}
                                for col, val in enumerate(values):
                                    headers.setdefault(val, col)
                            continue
                        def get(*names):
                            for name in names:
                                if name in headers:
                                    return row[headers[name]]
                            return None
                        code = text(get('编号'))
                        if not code:
                            if current_record is not None and any(v is not None for v in row):
                                current_record['catalog_continuation_rows'].append(dict(
                                    source_row=rownum, fields=catalog_cells(row, rownum, header_rows)))
                                current_record['catalog_reference_tests'].extend(
                                    catalog_tests(row, rownum, reference_start, header_rows[-1]))
                            continue
                        capacity = get('标称mah', '标称容量')
                        try:
                            capacity = float(capacity) / 1000 if capacity is not None else None
                        except (ValueError, TypeError):
                            capacity = None
                        purchase = get('购买日期')
                        purchase = purchase.date().isoformat() if isinstance(purchase, datetime) else text(purchase) or None
                        resistance_col = headers.get('交流电阻', headers.get('交流内阻'))
                        voltage_col = headers.get('内阻测量电压', headers.get('测量电压', headers.get('测试电压')))
                        if voltage_col is None and form_factor(path) == '21700' and resistance_col == 17:
                            voltage_col = 16  # Confirmed source layout: Q voltage / R AC resistance.
                        if voltage_col is None and Path(path).name == '其他电池记录.xlsx' and resistance_col == 19:
                            voltage_col = 18  # Same catalog layout: S voltage / T AC resistance.
                        if voltage_col is None and form_factor(path) == '软包' and resistance_col == 18:
                            voltage_col = 17  # Source layout: R voltage / S AC resistance.
                        reference_start = resistance_col + 1 if resistance_col is not None else len(row)
                        current_record = dict(form_factor=form_factor(path), code=code,
                            manufacturer=text(get('厂家')), supplier=text(get('购买渠道')),
                            condition=text(get('状态')), nominal_capacity_ah=capacity,
                            purchase_date=purchase, rated_c_rate=get('标称倍率','倍率'),
                            ac_resistance_mohm=row[resistance_col] if resistance_col is not None else None,
                            ac_resistance_voltage_v=row[voltage_col] if voltage_col is not None else None,
                            usage=text(get('用途')),
                            model_spec=text(get('型号')), dimensions=text(get('尺寸')), price=get('价格'),
                            nominal_voltage_v=get('标称电压'), weight_g=get('重量/g'),
                            ac_resistance_note='原表值；测试架比开尔文夹多0.4mΩ的备注未用于自动校正' if form_factor(path)=='21700' else '',
                            summary_reference=[text(v) if isinstance(v,(date,datetime)) else v for v in row[reference_start:]],
                            summary_reference_start_column=reference_start+1,
                            catalog_fields=catalog_cells(row, rownum, header_rows),
                            catalog_reference_tests=catalog_tests(row, rownum, reference_start, header_rows[-1]),
                            catalog_continuation_rows=[],
                            mapping_hash=source_hash,
                            appearance=text(get('外观')), mapping_source=str(Path(path).resolve()),
                            mapping_sheet=sheet.title, mapping_row=rownum)
                        records.append(current_record)
            finally:
                book.close()
        except Exception as exc:
            issues.append(f'汇总表读取失败 {Path(path).name}: {exc}')
    return records, issues


def filename_metadata(path):
    name = path.stem
    result = dict(form_factor=form_factor(path), category_code='', sample_code='',
                  test_date=None, filename_time=None, recorded_start=None,
                  set_current_a=None, cutoff_voltage_v=None, mode=None,
                  initial_soc_pct=100.0, reference_capacity_ah=None,
                  soc_basis='本次实测容量归一化；初始SOC=100%（用户默认）')
    match = re.match(r'(\d{4}-\d{2}-\d{2}-\d{2}-\d{2}-\d{2})', name)
    if match:
        try:
            result['filename_time'] = datetime.strptime(match[1], '%Y-%m-%d-%H-%M-%S').isoformat(' ')
        except ValueError:
            pass
    match = re.search(r'(?:21700|18650)-(\d+)-(\d+)(?=-|$)', name)
    if match:
        result.update(category_code=match[1], sample_code=match[2])
        result['summary_code'] = f'{match[1]}-{match[2]}'
    else:
        legacy = re.search(r'(?:21700|18650)编号([0-9]+)(Full)?$', name)
        if legacy:
            result.update(category_code=legacy[1], sample_code='原编号'+legacy[1], summary_code=legacy[1])
        else:
            legacy_current = re.search(r'(?:21700|18650)-([0-9]+)-(\d+(?:\.\d+)?)A$', name, re.I)
            if legacy_current:
                result.update(category_code=legacy_current[1], sample_code='原编号'+legacy_current[1], summary_code=legacy_current[1])
            elif not re.search(r'应该|可能|串联|电池组|21700-|18650-',name):
                pair = re.search(r'DL24-\d+-(\d+)-(\d+)(?=-|$)',name)
                if pair:
                    result.update(category_code=pair[1],sample_code=pair[2],summary_code=f'{pair[1]}-{pair[2]}')
    result['protocol_class'] = '非标准放电' if '非标准放电' in str(path) else '常规放电'
    for pattern, key in ((r'(?:^|-)(\d+(?:\.\d+)?)A(?=-|$)', 'set_current_a'),
                         (r'(?:^|-)(\d+(?:\.\d+)?)V(?=-|$)', 'cutoff_voltage_v')):
        match = re.search(pattern, name, re.I)
        if match:
            result[key] = float(match[1])
    result['condition_source'] = '文件名（标称条件）；实测电流与末端电压另存'
    c_rate = re.search(r'(?:^|-)(\d+(?:\.\d+)?)C(?=-|$)', name, re.I)
    result['test_c_rate'] = float(c_rate[1]) if c_rate else None
    return result


class CsvSheet:
    title = 'CSV'
    def __init__(self, rows):
        self.rows = rows
    def iter_rows(self, values_only=True):
        for row in self.rows:
            yield tuple(v if v != '' else None for v in row) + (None,) * max(0, 10-len(row))


@contextmanager
def open_source(path):
    if path.suffix.lower() == '.csv':
        for encoding in ('utf-8-sig','gb18030'):
            try:
                content = path.read_text(encoding=encoding)
                break
            except UnicodeError:
                continue
        else:
            raise ValueError('CSV 编码无法识别')
        yield [CsvSheet(list(csv.reader(content.splitlines()))) ]
    else:
        book = openpyxl.load_workbook(path,read_only=True,data_only=True)
        try:
            yield book
        finally:
            book.close()


def parse_time(value, header):
    if isinstance(value, datetime):
        return value.timestamp(), value.isoformat(' ')
    if isinstance(value, timedelta):
        return value.total_seconds(), None
    s = text(value)
    try:
        value = datetime.fromisoformat(s.replace('_', ' '))
        return value.timestamp(), value.isoformat(' ')
    except ValueError:
        pass
    unit = header.lower().replace(' ', '')
    factor = 0.001 if '毫秒' in unit or '(ms)' in unit else 60 if '分钟' in unit or '(min)' in unit else 3600 if '小时' in unit or '(h)' in unit else 1 if '秒' in unit or '(s)' in unit else None
    if factor is not None:
        return float(s) * factor, None
    raise ValueError('时间格式或数值时间单位不明确')


def _same_electrical_sample(left, right):
    """Whether adjacent rows repeat one electrical sample at one timestamp."""
    return (
        left.get('raw_time') == right.get('raw_time')
        and left.get('voltage_v') == right.get('voltage_v')
        and left.get('current_a') == right.get('current_a')
    )


def deduplicate_exact_rows(points):
    """Keep the first of adjacent repeated electrical samples.

    The source workbook/CSV is never changed.  The returned source row numbers
    make the normalization auditable in the imported test metadata.  A
    temperature-only difference is retained as a warning because temperature
    is auxiliary to the electrical curve and cannot justify a second time
    interval.
    """
    kept, removed, notes = [], [], []
    for point in points:
        if kept and kept[-1].get('raw_time') == point.get('raw_time'):
            previous = kept[-1]
            if _same_electrical_sample(previous, point):
                removed.append(point.get('source_row'))
                if previous.get('temperature_c') != point.get('temperature_c'):
                    notes.append(
                        f"重复时间戳的温度由 {previous.get('temperature_c')} 变为 {point.get('temperature_c')} ℃，"
                        f"保留源行 {previous.get('source_row')}"
                    )
            else:
                # A DL24 timestamp has one-second resolution.  When the same
                # timestamp carries a different measurement, the later row is
                # the user's selected representative for that instant.
                removed.append(previous.get('source_row'))
                notes.append(
                    f"同一时间戳测量值不同，按规则去除源行 {previous.get('source_row')}，"
                    f"保留后一个源行 {point.get('source_row')}"
                )
                kept[-1] = point
        else:
            kept.append(point)
    return kept, removed, notes


def drop_isolated_nonpositive_voltage_rows(points):
    """Drop isolated non-positive voltage glitches between valid samples."""
    kept, removed = [], []
    for index, point in enumerate(points):
        try:
            voltage = float(point.get('voltage_v'))
        except (TypeError, ValueError):
            voltage = None
        previous = points[index - 1] if index > 0 else None
        following = points[index + 1] if index + 1 < len(points) else None
        try:
            previous_voltage = float(previous.get('voltage_v')) if previous is not None else None
        except (TypeError, ValueError):
            previous_voltage = None
        try:
            following_voltage = float(following.get('voltage_v')) if following is not None else None
        except (TypeError, ValueError):
            following_voltage = None
        isolated = (
            voltage is not None and voltage <= 0
            and previous is not None and following is not None
            and previous_voltage is not None and previous_voltage > 0
            and following_voltage is not None and following_voltage > 0
        )
        if isolated:
            removed.append(point.get('source_row'))
        else:
            kept.append(point)
    return kept, removed


def _manual_source_row_exclusions(path, source_row_exclusions):
    """Resolve manually reviewed source-row exclusions for one source file."""
    if not source_row_exclusions:
        return set()
    resolved = str(Path(path).resolve())
    for key, rows in source_row_exclusions.items():
        try:
            if str(Path(key).resolve()) == resolved:
                return {int(row) for row in rows}
        except (TypeError, ValueError):
            continue
    return set()


def _manual_form_factor_override(path, form_factor_overrides):
    """Resolve a reviewed form-factor override for one source file."""
    if not form_factor_overrides:
        return None
    resolved = str(Path(path).resolve())
    for key, value in form_factor_overrides.items():
        try:
            if str(Path(key).resolve()) == resolved:
                return text(value) or None
        except (TypeError, ValueError):
            continue
    return None


def preview_workbooks(paths, mapping_paths=(), source_row_exclusions=None, form_factor_overrides=None):
    """Build read-only candidates, optionally applying reviewed row exclusions.

    ``source_row_exclusions`` maps a source path to source-row numbers that a
    reviewer explicitly approved for removal.  The original workbook/CSV is
    never changed; the removed rows are recorded in candidate metadata and
    warnings so the imported test remains auditable.

    ``form_factor_overrides`` maps a source path to a reviewer-confirmed
    form factor when the filename does not carry it.  The override is stored
    in candidate metadata and does not alter the source file.
    """
    mappings, mapping_issues = load_mappings(mapping_paths)
    candidates = []
    for filename in paths:
        path = Path(filename)
        try:
            digest = hashlib.sha256(path.read_bytes()).hexdigest()
            with open_source(path) as book:
                for sheet in book:
                    metadata = filename_metadata(path)
                    override = _manual_form_factor_override(path, form_factor_overrides)
                    if override:
                        metadata['form_factor'] = override
                        metadata['form_factor_override'] = override
                    c = dict(metadata=metadata, points=[], issues=list(mapping_issues),
                             source_file=str(path.resolve()), source_sheet=sheet.title,
                             source_hash=digest, approved=False, summary={}, warnings=[])
                    headers = None
                    times, modes = [], set()
                    for n, row in enumerate(sheet.iter_rows(values_only=True), 1):
                        if headers is None:
                            names = [text(v).replace('（', '(').replace('）', ')') for v in row]
                            if '电压(V)' in names and '电流(A)' in names:
                                ti = next((i for i, h in enumerate(names) if h.startswith('时间')), None)
                                if ti is not None:
                                    headers = (ti, names.index('电压(V)'), names.index('电流(A)'),
                                               names.index('温度(℃)') if '温度(℃)' in names else None,
                                               names.index('模式') if '模式' in names else None, names[ti])
                            continue
                        ti, vi, ii, temp, mi, th = headers
                        if all(row[k] is None for k in (ti, vi, ii)):
                            continue
                        p = dict(source_row=n, raw_time=text(row[ti]), voltage_v=row[vi],
                                 current_a=row[ii], temperature_c=row[temp] if temp is not None else None)
                        try:
                            t, iso = parse_time(row[ti], th)
                            if not c['points']:
                                c['metadata']['test_date'] = iso
                                c['metadata']['recorded_start'] = iso
                            times.append(t)
                            p['elapsed_s'] = t - times[0]
                        except (ValueError, TypeError, OverflowError):
                            p['elapsed_s'] = None
                            c['issues'].append(f'第 {n} 行时间无法解析，未跨越该行积分')
                        if mi is not None:
                            modes.add(text(row[mi]))
                        c['points'].append(p)
                    if headers is None:
                        c['issues'].append('未找到原始时间、电压(V)、电流(A)表头')
                    meta = c['metadata']
                    c['points'], removed_rows, duplicate_notes = deduplicate_exact_rows(c['points'])
                    if removed_rows:
                        meta['duplicate_rows_removed'] = removed_rows
                        c['warnings'].append(
                            f'发现重复时间戳，按审查规则保留一个采样并去除 {len(removed_rows)} 行：'
                            + ', '.join(str(row) for row in removed_rows)
                        )
                        c['warnings'].extend(duplicate_notes)
                    requested_rows = _manual_source_row_exclusions(path, source_row_exclusions)
                    if requested_rows:
                        retained = []
                        removed_manual = []
                        for point in c['points']:
                            if point.get('source_row') in requested_rows:
                                removed_manual.append(point.get('source_row'))
                            else:
                                retained.append(point)
                        missing_manual = sorted(requested_rows - set(removed_manual))
                        if removed_manual:
                            c['points'] = retained
                            meta['manual_rows_removed'] = removed_manual
                            c['warnings'].append(
                                '按人工审查去除源行：' + ', '.join(str(row) for row in removed_manual)
                            )
                        if missing_manual:
                            c['issues'].append(
                                '人工审查指定去除的源行未找到：' + ', '.join(str(row) for row in missing_manual)
                            )
                    c['points'], isolated_voltage_rows = drop_isolated_nonpositive_voltage_rows(c['points'])
                    if isolated_voltage_rows:
                        meta['isolated_voltage_rows_removed'] = isolated_voltage_rows
                        c['warnings'].append(
                            '按人工审查规则去除孤立非正电压源行：'
                            + ', '.join(str(row) for row in isolated_voltage_rows)
                        )
                    meta['mode'] = ','.join(sorted(modes)) or None
                    if len(modes) > 1:
                        c['issues'].append('存在多个仪器模式，需确认实验阶段边界')
                    key = meta.get('summary_code', f"{meta['category_code']}-{meta['sample_code']}")
                    exact = [m for m in mappings if m['form_factor'] == meta['form_factor'] and m['code'] == key]
                    # Pouch serial numbers are not the cylinder category/sample convention.
                    if meta['form_factor'] == '软包':
                        exact = [m for m in mappings if m['form_factor'] == '软包' and re.search(r'(?<!\d)' + re.escape(m['code']) + r'(?!\d)', path.stem)]
                        if len(exact) == 1:
                            meta['category_code'], meta['sample_code'] = catalog_identity(exact[0])
                    if len(exact) == 1:
                        meta.update({k: v for k, v in exact[0].items() if k not in ('code', 'form_factor')})
                        meta['summary_code'] = exact[0]['code']
                        meta['model_name'] = ' '.join(filter(None, [meta.get('manufacturer'),meta.get('model_spec'),
                            f"{meta['nominal_capacity_ah']*1000:g}mAh" if meta.get('nominal_capacity_ah') else '', meta.get('condition')]))
                        if meta.get('test_c_rate') is not None and meta.get('nominal_capacity_ah'):
                            meta['set_current_a'] = meta['test_c_rate'] * meta['nominal_capacity_ah']
                            meta['condition_source'] = '文件名倍率 × 汇总表标称容量；截止设定未知时保留为空'
                    else:
                        c['issues'].append('汇总表未唯一匹配完整样品编号，需人工核定；未套用类别行')
                    if not meta['category_code'] or not meta['sample_code']:
                        c['issues'].append('文件名中的类别或样品编号不明确')
                    if not meta['test_date']:
                        c['issues'].append('缺少第一条原始记录的完整日期时间')
                    # Reviewed policy: preserve genuine sampling gaps and
                    # integrate over them when they are no longer than 100 s.
                    # Longer gaps remain review-blocking anomalies.
                    gap = 100.0
                    meta['max_gap_s'] = gap
                    points, summary, issues = calculate_points(c['points'], max_gap_s=gap)
                    c['issues'].extend(issues)
                    if summary['integration_valid'] and not any(p['elapsed_s'] is None for p in points):
                        meta['reference_capacity_ah'] = summary['total_capacity_ah']
                        points, summary, _ = calculate_points(points, 100.0, meta['reference_capacity_ah'], meta['soc_basis'], gap)
                    c.update(points=points, summary=summary)
                    currents = [p['current_a'] for p in points if p['current_a'] is not None and p['current_a'] > 0]
                    meta['measured_current_a'] = statistics.median(currents) if currents else None
                    meta['last_voltage_v'] = points[-1]['voltage_v'] if points else None
                    if meta['set_current_a'] is None or meta['cutoff_voltage_v'] is None:
                        c['warnings'].append('部分设定条件未记录，保留为空；实测中位电流/末端电压另存，不冒充设定值')
                    meta['import_warnings'] = c['warnings']
                    raw = [[p.get(k) for k in ('raw_time','voltage_v','current_a')] for p in points]
                    c['fingerprint'] = hashlib.sha256(json.dumps([meta['form_factor'], meta['category_code'], meta['sample_code'], raw], ensure_ascii=False, default=str).encode()).hexdigest()
                    candidates.append(c)
        except Exception as exc:
            candidates.append(dict(metadata=filename_metadata(path), points=[], issues=[f'文件读取失败：{exc}'],
                summary={}, source_file=str(path.resolve()), source_sheet='', source_hash='', fingerprint='', approved=False))
    return candidates
